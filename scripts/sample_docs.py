"""Builders for the sample corpus (and test fixtures).

Generates a small knowledge base for the fictional company "Northwind Cloud" in all
supported formats: Markdown, HTML, DOCX, PDF (multi-page) and TXT.

    python -m scripts.sample_docs            # writes ./samples/
"""

from __future__ import annotations

import io
import textwrap
from pathlib import Path


def make_pdf(pages: list[str], title: str | None = None) -> bytes:
    """Minimal, dependency-free PDF writer (Helvetica text, one string per page)."""

    def esc(s: str) -> str:
        return s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    objects: dict[int, str] = {}
    kids: list[int] = []
    next_id = 5
    for text in pages:
        lines: list[str] = []
        for para in text.split("\n"):
            lines.extend(textwrap.wrap(para, 90) or [""])
        stream = "BT /F1 11 Tf 50 750 Td 14 TL " + " ".join(f"({esc(ln)}) Tj T*" for ln in lines) + " ET"
        content_id, page_id = next_id, next_id + 1
        next_id += 2
        objects[content_id] = f"<< /Length {len(stream.encode('latin-1'))} >>\nstream\n{stream}\nendstream"
        objects[page_id] = (
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Contents {content_id} 0 R /Resources << /Font << /F1 3 0 R >> >> >>"
        )
        kids.append(page_id)
    objects[1] = "<< /Type /Catalog /Pages 2 0 R >>"
    objects[2] = f"<< /Type /Pages /Kids [{' '.join(f'{k} 0 R' for k in kids)}] /Count {len(kids)} >>"
    objects[3] = "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"
    objects[4] = f"<< /Title ({esc(title or '')}) /Producer (hybrid-rag samples) >>"

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets: dict[int, int] = {}
    for oid in sorted(objects):
        offsets[oid] = out.tell()
        out.write(f"{oid} 0 obj\n{objects[oid]}\nendobj\n".encode("latin-1"))
    xref = out.tell()
    size = max(objects) + 1
    out.write(f"xref\n0 {size}\n0000000000 65535 f \n".encode())
    for oid in range(1, size):
        out.write(f"{offsets[oid]:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {size} /Root 1 0 R /Info 4 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def make_docx(blocks: list[tuple[str, object]], title: str | None = None) -> bytes:
    """blocks: ("h1"|"h2"|"p"|"bullet", text) or ("table", [[cells], ...])."""
    import docx

    d = docx.Document()
    if title:
        d.core_properties.title = title
    for kind, value in blocks:
        if kind in ("h1", "h2", "h3"):
            d.add_heading(str(value), level=int(kind[1]))
        elif kind == "bullet":
            d.add_paragraph(str(value), style="List Bullet")
        elif kind == "table":
            rows = value  # type: ignore[assignment]
            table = d.add_table(rows=len(rows), cols=len(rows[0]))  # type: ignore[arg-type]
            for r, row in enumerate(rows):  # type: ignore[arg-type]
                for c, cell in enumerate(row):
                    table.cell(r, c).text = str(cell)
        else:
            d.add_paragraph(str(value))
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


REFUND_POLICY_MD = """# Northwind Cloud Refund Policy

Northwind Cloud wants every customer to be satisfied with their subscription. This policy explains
when refunds are available and how to request one.

## Eligibility

Monthly subscriptions can be refunded in full if the refund is requested within 30 days of the
initial purchase. Renewals of monthly subscriptions are not refundable, but you can cancel at any
time to stop the next renewal.

Annual subscriptions cancelled within the first 30 days receive a full refund. After 30 days,
annual subscriptions receive a prorated refund for the unused full months remaining in the term,
minus a 10% early-termination fee.

## Exceptions

Usage-based charges (compute hours, egress bandwidth and API overage) are never refundable,
because the resources have already been consumed. Accounts suspended for violating the Acceptable
Use Policy are not eligible for any refund.

## How to request a refund

Open the Billing page in the Northwind console and choose "Request refund", or email
billing@northwind.example with your account ID and invoice number. Refunds are issued to the
original payment method within 5 to 10 business days after approval.
"""

SECURITY_HANDBOOK_HTML = """<!doctype html>
<html><head><title>Northwind Security Handbook</title>
<script>console.log("analytics");</script></head>
<body>
<nav><a href="/">Home</a> | <a href="/docs">Docs</a></nav>
<main>
<h1>Northwind Security Handbook</h1>
<p>This handbook defines the security practices every Northwind employee and contractor must follow.</p>
<h2>Passwords and Multi-Factor Authentication</h2>
<p>Passwords must be at least 14 characters long and must not be reused across services.
Use the company password manager to generate and store credentials.</p>
<p>Multi-factor authentication (MFA) is mandatory for all production systems, the cloud console
and email. Hardware security keys are required for administrators.</p>
<h2>Incident Reporting</h2>
<p>Any suspected security incident must be reported to the security team within 24 hours by
emailing security@northwind.example or paging the on-call engineer.</p>
<ul><li>Lost or stolen laptops must be reported within 4 hours.</li>
<li>Phishing emails should be forwarded to phishing@northwind.example.</li></ul>
<h2>Data Retention</h2>
<p>Customer data is retained for 90 days after account deletion and then permanently erased
from all primary storage. Encrypted backups are kept for a further 35 days.</p>
</main>
<footer>Copyright Northwind Cloud</footer>
</body></html>
"""

PRODUCT_FAQ_TXT = """Northwind Cloud Product FAQ

What are the API rate limits?
The public REST API allows 1,000 requests per minute per API key on the Standard plan and 5,000
requests per minute on the Enterprise plan. Requests over the limit receive HTTP 429 responses.

Which regions are available?
Northwind Cloud runs in three regions: eu-central (Frankfurt), us-east (Virginia) and ap-south
(Singapore). Data never leaves the region selected when the workspace is created.

Can I export my data?
Yes. Workspace owners can export all data as Parquet or JSON from Settings > Export at any time.
Exports larger than 50 GB are delivered to a customer-owned object storage bucket.
"""

SLA_PAGES = [
    "Northwind Cloud Service Level Agreement\n"
    "1. Uptime Commitment\n"
    "Northwind Cloud commits to a monthly uptime of 99.95% for the Compute and Storage services. "
    "Uptime is measured by external probes every minute from three independent locations. "
    "Scheduled maintenance announced at least 7 days in advance is excluded from the calculation.",
    "2. Service Credits\n"
    "If monthly uptime falls below 99.95%, customers receive a service credit of 10% of the monthly fee. "
    "If uptime falls below 99.0%, the credit increases to 25% of the monthly fee. "
    "If uptime falls below 95.0%, the credit is 50% of the monthly fee. "
    "Service credits must be requested within 30 days of the end of the affected month.",
    "3. Support Response Times\n"
    "Severity 1 incidents (production down) receive a first response within 30 minutes, 24 hours a day. "
    "Severity 2 incidents (degraded service) receive a first response within 4 business hours. "
    "Severity 3 questions receive a first response within 2 business days.",
]

EMPLOYEE_HANDBOOK_BLOCKS: list[tuple[str, object]] = [
    ("h1", "Northwind Employee Handbook"),
    ("p", "Welcome to Northwind Cloud. This handbook summarizes our core people policies."),
    ("h2", "Paid Time Off"),
    (
        "p",
        "Full-time employees receive 25 days of paid vacation per calendar year, plus public holidays. "
        "Up to 5 unused days can be carried over to the next year.",
    ),
    ("h2", "Remote Work"),
    (
        "p",
        "Employees may work remotely up to 3 days per week. Fully remote arrangements require approval "
        "from both the team manager and the People team.",
    ),
    ("bullet", "Core collaboration hours are 10:00 to 15:00 in the team's home time zone."),
    ("bullet", "A home-office stipend of 500 EUR is available once every two years."),
    ("h2", "Expense Reimbursement"),
    (
        "p",
        "Business expenses are reimbursed within 14 days of an approved expense report. "
        "Daily limits per category are listed below.",
    ),
    (
        "table",
        [
            ["Category", "Daily limit"],
            ["Meals", "60 EUR"],
            ["Hotel", "180 EUR"],
            ["Local transport", "40 EUR"],
        ],
    ),
]


def sample_files() -> dict[str, bytes]:
    return {
        "refund_policy.md": REFUND_POLICY_MD.encode(),
        "security_handbook.html": SECURITY_HANDBOOK_HTML.encode(),
        "product_faq.txt": PRODUCT_FAQ_TXT.encode(),
        "sla_agreement.pdf": make_pdf(SLA_PAGES, title="Northwind Cloud Service Level Agreement"),
        "employee_handbook.docx": make_docx(EMPLOYEE_HANDBOOK_BLOCKS, title="Northwind Employee Handbook"),
    }


def main() -> None:
    out = Path("samples")
    out.mkdir(exist_ok=True)
    for name, data in sample_files().items():
        (out / name).write_bytes(data)
        print(f"wrote samples/{name} ({len(data)} bytes)")


if __name__ == "__main__":
    main()
